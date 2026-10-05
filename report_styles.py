import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Color palette
COLOR_PRIMARY = RGBColor(30, 58, 138)     # Navy Blue #1E3A8A
COLOR_SECONDARY = RGBColor(14, 116, 144)  # Ocean Cyan #0E7490
COLOR_TEXT = RGBColor(31, 41, 55)         # Charcoal Dark Gray #1F2937
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Gray #64748B
HEX_PRIMARY = "1E3A8A"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_BOX_BG = "F1F5F9"
HEX_ACCENT_LEFT = "2563EB"

def setup_document():
    doc = docx.Document()
    
    # Page setup: Standard academic margins (Left 3cm, Top 2cm, Right 2cm, Bottom 2cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)      # 2.0 cm
        section.bottom_margin = Inches(0.79)   # 2.0 cm
        section.left_margin = Inches(1.18)     # 3.0 cm
        section.right_margin = Inches(0.79)    # 2.0 cm
        section.different_first_page_header_footer = True
        
        # Header / Footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = p_ft.add_run("Báo cáo môn học: Kiến trúc và Thiết kế Phần mềm")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(9)
        f_run.font.italic = True
        f_run.font.color.rgb = COLOR_MUTED

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = COLOR_TEXT
    normal_style.paragraph_format.line_spacing = 1.3
    normal_style.paragraph_format.space_after = Pt(6)
    
    return doc

def add_p(doc, text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, font_size=13):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(font_size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = COLOR_TEXT
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = COLOR_SECONDARY
    return p

def add_heading_4(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    r.italic = True
    r.font.color.rgb = COLOR_TEXT
    return p

def add_bullet(doc, text, bold_prefix="", level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(13)
        r_pre.bold = True
        r_pre.font.color.rgb = COLOR_TEXT
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_TEXT
    return p

def add_diagram_box(doc, ascii_art, caption=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.2)
    
    # Border & background shading
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>\n'
        f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_ACCENT_LEFT}"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_BOX_BG}"/>')
    tcPr.append(shd)
    
    # Text inside box
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    
    r = p.add_run(ascii_art)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(15, 23, 42)
    
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(11)
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_MUTED

def add_table_styled(doc, headers, data, col_widths=None, caption=""):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(11.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = COLOR_PRIMARY

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Style Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        cell = hdr_cells[i]
        if col_widths and i < len(col_widths):
            cell.width = col_widths[i]
        cell_p = cell.paragraphs[0]
        cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell_p.paragraph_format.line_spacing = 1.15
        cell_p.paragraph_format.space_before = Pt(5)
        cell_p.paragraph_format.space_after = Pt(5)
        r = cell_p.add_run(h_text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_PRIMARY}"/>')
        tcPr.append(shd)
        
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY}"/>\n'
            f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>\n'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY}"/>\n'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

    # Fill data rows
    for row_idx, row_values in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_fill = HEX_LIGHT_BG if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_values):
            cell = row_cells[col_idx]
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]
            cell_p = cell.paragraphs[0]
            cell_p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            cell_p.paragraph_format.line_spacing = 1.15
            cell_p.paragraph_format.space_before = Pt(4)
            cell_p.paragraph_format.space_after = Pt(4)
            r = cell_p.add_run(str(cell_value))
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.color.rgb = COLOR_TEXT
            
            tcPr = cell._tc.get_or_add_tcPr()
            if bg_fill != "FFFFFF":
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_fill}"/>')
                tcPr.append(shd)
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>\n'
                f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>\n'
                f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>\n'
                f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>\n'
                f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>\n'
                f'</w:tcBorders>'
            )
            tcPr.append(borders)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table
