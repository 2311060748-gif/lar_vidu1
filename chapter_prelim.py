import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_styles import (
    add_p, add_heading_1, add_heading_2, add_bullet, 
    add_table_styled, COLOR_PRIMARY, COLOR_SECONDARY, COLOR_TEXT
)

def create_cover_page(doc, is_sub=False):
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_before = Pt(10)
    p_top.paragraph_format.space_after = Pt(2)
    r1 = p_top.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\nKHOA CÔNG NGHỆ THÔNG TIN")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_TEXT
    
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(36)
    r_l = p_line.add_run("--------------------***--------------------")
    r_l.font.bold = True
    
    p_mid = doc.add_paragraph()
    p_mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mid.paragraph_format.space_after = Pt(12)
    r_sub = p_mid.add_run("BÁO CÁO BÀI TẬP LỚN / TIỂU LUẬN MÔN HỌC" if not is_sub else "TRANG BÌA PHỤ - BÁO CÁO MÔN HỌC")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(15)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SECONDARY
    
    p_subj = doc.add_paragraph()
    p_subj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subj.paragraph_format.space_after = Pt(20)
    r_subj = p_subj.add_run("HỌC PHẦN: KIẾN TRÚC VÀ THIẾT KẾ PHẦN MỀM")
    r_subj.font.name = "Times New Roman"
    r_subj.font.size = Pt(16)
    r_subj.font.bold = True
    r_subj.font.color.rgb = COLOR_PRIMARY
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(36)
    r_t1 = p_title.add_run("ĐỀ TÀI:\n")
    r_t1.font.name = "Times New Roman"
    r_t1.font.size = Pt(15)
    r_t1.font.bold = True
    
    r_t2 = p_title.add_run("NGHIÊN CỨU KIẾN TRÚC, NGUYÊN LÝ THIẾT KẾ VÀ XÂY DỰNG\nHỆ THỐNG THƯƠNG MẠI ĐIỆN TỬ VÀ QUẢN LÝ CỬA HÀNG PHỤ TÙNG XE MÁY TRỰC TUYẾN\n(ONLINE MOTORCYCLE SPARE PARTS STORE & INVENTORY SYSTEM - MOTOPARTS)")
    r_t2.font.name = "Times New Roman"
    r_t2.font.size = Pt(16)
    r_t2.font.bold = True
    r_t2.font.color.rgb = COLOR_PRIMARY
    
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Inches(1.5)
    p_info.paragraph_format.space_after = Pt(50)
    p_info.paragraph_format.line_spacing = 1.3
    
    r_info = p_info.add_run(
        "Giảng viên hướng dẫn : TS. Nguyễn Văn A\n"
        "Bộ môn              : Kỹ thuật Phần mềm\n"
        "Sinh viên thực hiện : Nhóm Nghiên cứu & Phát triển\n"
        "Mã số sinh viên     : 20210001 - 20210002\n"
        "Lớp                 : KTPM - K16\n"
        "Hệ đào tạo          : Đại học chính quy"
    )
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(13)
    r_info.font.color.rgb = COLOR_TEXT
    
    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bot.paragraph_format.space_before = Pt(30)
    r_bot = p_bot.add_run("Hà Nội, Năm học 2025 - 2026")
    r_bot.font.name = "Times New Roman"
    r_bot.font.size = Pt(13)
    r_bot.font.bold = True
    
    doc.add_page_break()

def add_acknowledgements(doc):
    add_heading_1(doc, "LỜI CẢM ƠN")
    add_p(doc, 
        "Lời đầu tiên, nhóm tác giả xin được bày tỏ lòng biết ơn sâu sắc và chân thành nhất tới Ban Giám hiệu, "
        "các thầy cô giáo Khoa Công nghệ Thông tin - Trường Đại học Công nghệ Thông tin và Truyền thông đã tận tình "
        "truyền đạt những kiến thức quý báu, phương pháp luận khoa học và tạo mọi điều kiện học tập, nghiên cứu thuận lợi nhất "
        "cho chúng em trong suốt quá trình theo học môn học Kiến trúc và Thiết kế Phần mềm."
    )
    add_p(doc,
        "Đặc biệt, nhóm chúng em xin gửi lời tri ân sâu sắc nhất tới Thầy/Cô giảng viên phụ trách học phần. Bằng sự tâm huyết, "
        "kiến thức chuyên sâu và những chỉ dẫn thực tiễn, Thầy/Cô đã định hướng tư duy thiết kế hệ thống, phân tích kiến trúc phần mềm, "
        "giúp chúng em không chỉ nắm vững cơ sở lý thuyết mà còn biết cách áp dụng các nguyên lý thiết kế SOLID, các mẫu thiết kế GoF, "
        "và các phong cách kiến trúc hiện đại vào việc xây dựng một hệ thống website thương mại điện tử chuyên ngành thực tế: "
        "'Hệ thống Thương mại Điện tử và Quản lý Cửa hàng Phụ tùng Xe máy Trực tuyến (MotoParts)'."
    )
    add_p(doc,
        "Trong quá trình thực hiện bài tập lớn, dù đã nỗ lực hết mình để khảo sát, phân tích, mô hình hóa và hiện thực hóa đề tài, "
        "song do kiến thức và kinh nghiệm thực tiễn còn những hạn chế nhất định, báo cáo chắc chắn khó tránh khỏi những thiếu sót. "
        "Nhóm chúng em rất mong nhận được những nhận xét, góp ý và đánh giá quý báu từ quý Thầy/Cô để bản báo cáo cũng như kiến thức "
        "chuyên môn của chúng em được hoàn thiện hơn nữa trong tương lai."
    )
    add_p(doc, "Chúng em xin chân thành cảm ơn!", italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    doc.add_page_break()

def add_table_of_contents(doc):
    add_heading_1(doc, "MỤC LỤC")
    
    toc_data = [
        ("Trang bìa chính", "i"),
        ("Trang bìa phụ", "ii"),
        ("Lời cảm ơn", "iii"),
        ("Mục lục", "iv"),
        ("Danh mục chữ viết tắt", "vi"),
        ("Danh mục bảng biểu", "vii"),
        ("Danh mục hình ảnh", "viii"),
        ("MỞ ĐẦU", "1"),
        ("  1. Lý do chọn đề tài", "1"),
        ("  2. Mục tiêu nghiên cứu", "2"),
        ("  3. Đối tượng và phạm vi nghiên cứu", "2"),
        ("  4. Phương pháp thực hiện", "3"),
        ("  5. Cấu trúc báo cáo", "3"),
        ("CHƯƠNG 1: TỔNG QUAN VỀ KIẾN TRÚC VÀ THIẾT KẾ PHẦN MỀM", "5"),
        ("  1.1. Thiết kế phần mềm trong quy trình phát triển", "5"),
        ("    1.1.1. Một số khái niệm", "5"),
        ("    1.1.2. Chuyển đổi sang mô hình thiết kế", "7"),
        ("  1.2. Kiến trúc phần mềm", "9"),
        ("    1.2.1. Khái niệm kiến trúc phần mềm", "9"),
        ("    1.2.2. Kiến trúc và các thuộc tính chất lượng", "11"),
        ("    1.2.3. Yêu cầu có ý nghĩa kiến trúc (ASRs)", "13"),
        ("  1.3. Các nguyên lý thiết kế", "15"),
        ("    1.3.1. Trừu tượng hóa và che giấu thông tin", "15"),
        ("    1.3.2. Phân tách mối quan tâm và mô đun hóa", "17"),
        ("    1.3.3. Tính gắn kết và tính ghép nối", "18"),
        ("    1.3.4. Nguyên lý SOLID", "21"),
        ("  1.4. Thể loại và phong cách kiến trúc", "26"),
        ("    1.4.1. Thể loại kiến trúc", "26"),
        ("    1.4.2. Phong cách kiến trúc", "28"),
        ("    1.4.3. So sánh thể loại và phong cách kiến trúc", "31"),
        ("  1.5. Mẫu thiết kế (Design Patterns)", "33"),
        ("  1.6. Kết luận chương 1", "37"),
        ("CHƯƠNG 2: PHÂN TÍCH YÊU CẦU VÀ THIẾT KẾ MÔ HÌNH", "39"),
        ("  2.1. Tổng quan bài toán", "39"),
        ("    2.1.1. Bối cảnh và hiện trạng thị trường phụ tùng xe máy", "39"),
        ("    2.1.2. Vấn đề cần giải quyết và mục tiêu hệ thống", "40"),
        ("    2.1.3. Phạm vi hệ thống", "42"),
        ("  2.2. Yêu cầu chức năng", "43"),
        ("    2.2.1. Danh mục yêu cầu chức năng", "43"),
        ("    2.2.2. Sơ đồ use case", "46"),
        ("    2.2.3. Đặc tả các use case trọng tâm", "48"),
        ("  2.3. Mô hình hóa quy trình nghiệp vụ", "55"),
        ("    2.3.1. Quy trình Đặt mua phụ tùng và thanh toán trực tuyến", "55"),
        ("    2.3.2. Quy trình Quản lý kho phụ tùng và cập nhật tồn kho", "58"),
        ("    2.3.3. Quy trình Tư vấn kỹ thuật phụ tùng và xác thực tài khoản", "60"),
        ("  2.4. Yêu cầu phi chức năng và ràng buộc", "62"),
        ("    2.4.1. Yêu cầu phi chức năng", "62"),
        ("    2.4.2. Ràng buộc hệ thống", "65"),
        ("  2.5. Kết luận chương 2", "66"),
        ("CHƯƠNG 3: THIẾT KẾ KIẾN TRÚC", "68"),
        ("  3.1. Yêu cầu và tiêu chí lựa chọn kiến trúc", "68"),
        ("  3.2. Lựa chọn phong cách kiến trúc", "70"),
        ("    3.2.1. Các phương án kiến trúc", "70"),
        ("    3.2.2. Quyết định lựa chọn kiến trúc", "73"),
        ("    3.2.3. Quan hệ giữa phong cách, mẫu kiến trúc và mẫu thiết kế", "75"),
        ("  3.3. Kiến trúc tổng thể của hệ thống", "77"),
        ("    3.3.1. Sơ đồ kiến trúc tổng thể", "77"),
        ("    3.3.2. Trách nhiệm của từng tầng", "79"),
        ("    3.3.3. Quy tắc phụ thuộc và trao đổi dữ liệu giữa các tầng", "82"),
        ("  3.4. Phân rã hệ thống thành các thành phần", "84"),
        ("    3.4.1. Nguyên tắc phân rã thành phần", "84"),
        ("    3.4.2. Danh mục thành phần và trách nhiệm", "85"),
        ("    3.4.3. Sơ đồ phân rã thành phần", "87"),
        ("    3.4.4. Quan hệ giữa các thành phần", "89"),
        ("  3.5. Đánh giá kiến trúc", "91"),
        ("    3.5.1. Phương pháp đánh giá (ATAM)", "91"),
        ("    3.5.2. Đánh giá theo các kịch bản chất lượng", "93"),
        ("    3.5.3. Điểm nhạy cảm, điểm đánh đổi và rủi ro", "96"),
        ("  3.6. Kết luận chương 3", "98"),
        ("CHƯƠNG 4: THIẾT KẾ DỮ LIỆU VÀ LỚP", "100"),
        ("  4.1. Thiết kế dữ liệu", "100"),
        ("    4.1.1. Thực thể, thuộc tính", "100"),
        ("    4.1.2. Mối quan hệ", "104"),
        ("    4.1.3. Sơ đồ thực thể quan hệ (ERD)", "107"),
        ("  4.2. Thiết kế lớp", "109"),
        ("    4.2.1. Xác định lớp và trách nhiệm", "109"),
        ("    4.2.2. Thuộc tính, phương thức và quan hệ giữa các lớp", "111"),
        ("    4.2.3. Sơ đồ lớp (Class Diagram)", "114"),
        ("  4.3. Sơ đồ trạng thái", "116"),
        ("    4.3.1. Sơ đồ trạng thái Đơn hàng phụ tùng (Order)", "116"),
        ("    4.3.2. Sơ đồ trạng thái Giao dịch thanh toán MoMo", "118"),
        ("    4.3.3. Sơ đồ trạng thái Tài khoản người dùng & OTP", "120"),
        ("  4.4. Kết luận chương 4", "122"),
        ("CHƯƠNG 5: THIẾT KẾ GIAO DIỆN VÀ THÀNH PHẦN", "124"),
        ("  5.1. Thiết kế giao diện", "124"),
        ("    5.1.1. Nguyên tắc thiết kế giao diện", "124"),
        ("    5.1.2. Người dùng và cấu trúc điều hướng", "127"),
        ("    5.1.3. Thiết kế các màn hình", "129"),
        ("    5.1.4. Thiết kế giao diện với hệ thống ngoài (MoMo, GHN, SMTP)", "133"),
        ("  5.2. Thiết kế thành phần", "136"),
        ("    5.2.1. Phạm vi và mục tiêu thiết kế thành phần", "136"),
        ("    5.2.2. Interface của các thành phần", "137"),
        ("    5.2.3. Thiết kế tương tác giữa các thành phần (Sơ đồ tuần tự)", "140"),
        ("    5.2.4. Thiết kế chi tiết các thành phần trọng tâm", "144"),
        ("    5.2.5. Đánh giá thiết kế thành phần", "147"),
        ("  5.3. Kết luận chương 5", "149"),
        ("KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", "151"),
        ("  1. Kết luận", "151"),
        ("  2. Hướng phát triển", "153"),
        ("TÀI LIỆU THAM KHẢO", "155")
    ]
    
    headers = ["Nội dung / Đề mục", "Trang"]
    col_widths = [Inches(5.2), Inches(1.0)]
    add_table_styled(doc, headers, toc_data, col_widths)
    doc.add_page_break()

def add_acronyms_tables_figures(doc):
    add_heading_1(doc, "DANH MỤC CHỮ VIẾT TẮT")
    acronyms = [
        ("ACID", "Atomicity, Consistency, Isolation, Durability", "Các thuộc tính đảm bảo tính tin cậy tuyệt đối của giao dịch cơ sở dữ liệu"),
        ("API", "Application Programming Interface", "Giao diện lập trình ứng dụng"),
        ("ASR", "Architecturally Significant Requirement", "Yêu cầu có ý nghĩa kiến trúc"),
        ("ATAM", "Architecture Tradeoff Analysis Method", "Phương pháp phân tích đánh đổi kiến trúc phần mềm"),
        ("COD", "Cash On Delivery", "Phương thức thanh toán tiền mặt trực tiếp khi nhận hàng"),
        ("CRUD", "Create, Read, Update, Delete", "Các thao tác cơ bản trên dữ liệu: Thêm, Đọc, Sửa, Xóa"),
        ("DTO", "Data Transfer Object", "Đối tượng truyền dữ liệu giữa các tầng kiến trúc"),
        ("ERD", "Entity Relationship Diagram", "Sơ đồ thực thể - quan hệ"),
        ("GHN", "Giao Hàng Nhanh", "Hệ thống dịch vụ bưu chính chuyển phát và logistics tại Việt Nam"),
        ("GoF", "Gang of Four", "Nhóm 4 tác giả khai sinh 23 mẫu thiết kế hướng đối tượng kinh điển"),
        ("HMAC", "Hash-based Message Authentication Code", "Mã xác thực thông điệp dựa trên hàm băm mật mã an toàn"),
        ("IPN", "Instant Payment Notification", "Cơ chế thông báo thanh toán tức thời từ máy chủ cổng thanh toán"),
        ("MVC", "Model - View - Controller", "Mẫu kiến trúc phân tách Dữ liệu - Giao diện - Bộ điều khiển"),
        ("OEM", "Original Equipment Manufacturer", "Nhà sản xuất phụ tùng và thiết bị gốc (Honda, Yamaha, DID...)"),
        ("ORM", "Object-Relational Mapping", "Kỹ thuật ánh xạ giữa đối tượng hướng đối tượng và CSDL quan hệ"),
        ("OTP", "One-Time Password", "Mật khẩu xác thực dùng một lần qua Email/SMS"),
        ("RBAC", "Role-Based Access Control", "Kiểm soát truy cập và phân quyền dựa trên vai trò người dùng"),
        ("REST", "Representational State Transfer", "Kiểu kiến trúc thiết kế dịch vụ web chuẩn HTTP"),
        ("SDLC", "Software Development Life Cycle", "Chu trình phát triển phần mềm"),
        ("SKU", "Stock Keeping Unit", "Mã phân loại phụ tùng tồn kho để quản lý sản phẩm"),
        ("SOLID", "SRP, OCP, LSP, ISP, DIP", "5 nguyên lý thiết kế hướng đối tượng nền tảng"),
        ("UI/UX", "User Interface / User Experience", "Giao diện người dùng / Trải nghiệm người dùng")
    ]
    add_table_styled(doc, ["Ký hiệu viết tắt", "Cụm từ tiếng Anh nguyên văn", "Ý nghĩa / Diễn giải kỹ thuật"], acronyms, [Inches(1.2), Inches(2.6), Inches(2.4)])
    
    add_heading_1(doc, "DANH MỤC BẢNG BIỂU")
    tables_list = [
        ("Bảng 1.1", "Bảng phân loại các mức độ gắn kết (Cohesion) và ghép nối (Coupling)"),
        ("Bảng 1.2", "Bảng so sánh chi tiết giữa Thể loại kiến trúc và Phong cách kiến trúc"),
        ("Bảng 1.3", "Bảng tổng hợp 23 mẫu thiết kế GoF phân loại theo 3 nhóm và ứng dụng trong hệ thống phụ tùng"),
        ("Bảng 2.1", "Bảng danh mục yêu cầu chức năng hệ thống bán phụ tùng (Guest, Customer, Admin)"),
        ("Bảng 2.2", "Bảng đặc tả ca sử dụng UC01: Tra cứu và Đặt mua phụ tùng trực tuyến"),
        ("Bảng 2.3", "Bảng đặc tả ca sử dụng UC02: Thanh toán đơn hàng phụ tùng qua cổng MoMo"),
        ("Bảng 2.4", "Bảng đặc tả ca sử dụng UC03: Quản lý danh mục phụ tùng xe máy và số lượng tồn kho"),
        ("Bảng 2.5", "Bảng đặc tả ca sử dụng UC04: Tư vấn kỹ thuật phụ tùng trực tuyến (Live Chat)"),
        ("Bảng 2.6", "Bảng tổng hợp yêu cầu phi chức năng hệ thống MotoParts theo mô hình FURPS+"),
        ("Bảng 3.1", "Ma trận quyết định lựa chọn phong cách kiến trúc hệ thống MotoParts"),
        ("Bảng 3.2", "Bảng danh mục các thành phần kiến trúc và phân công trách nhiệm"),
        ("Bảng 3.3", "Bảng đánh giá kiến trúc theo phương pháp ATAM qua các kịch bản chất lượng"),
        ("Bảng 4.1", "Từ điển dữ liệu bảng Users (Tài khoản người dùng, Thợ kỹ thuật & Admin)"),
        ("Bảng 4.2", "Từ điển dữ liệu bảng Categories (Danh mục phụ tùng: Má phanh, Lọc gió, Bugi, Nhông xích...)"),
        ("Bảng 4.3", "Từ điển dữ liệu bảng Products (Chi tiết phụ tùng: Mã OEM, Tên, Giá, Tồn kho, Hình ảnh)"),
        ("Bảng 4.4", "Từ điển dữ liệu bảng Orders (Đơn đặt hàng phụ tùng, địa chỉ, trạng thái GHN, tiền thanh toán)"),
        ("Bảng 4.5", "Từ điển dữ liệu bảng OrderItems (Chi tiết số lượng và đơn giá từng phụ tùng trong đơn)"),
        ("Bảng 4.6", "Từ điển dữ liệu bảng PaymentTransactions (Lịch sử giao dịch thanh toán MoMo)"),
        ("Bảng 4.7", "Từ điển dữ liệu bảng Messages (Tin nhắn tư vấn kỹ thuật phụ tùng trực tuyến)"),
        ("Bảng 4.8", "Bảng mô tả trách nhiệm các lớp chính trong hệ thống phần mềm"),
        ("Bảng 5.1", "Bảng tham chiếu 10 nguyên lý Usability Heuristics của Nielsen áp dụng vào giao diện phụ tùng"),
        ("Bảng 5.2", "Bảng mô tả cấu trúc API tương tác với Cổng thanh toán MoMo"),
        ("Bảng 5.3", "Bảng đánh giá đo lường chất lượng thiết kế thành phần hệ thống phụ tùng")
    ]
    add_table_styled(doc, ["Ký hiệu bảng", "Tên bảng biểu mô tả chi tiết"], tables_list, [Inches(1.5), Inches(4.7)])
    
    add_heading_1(doc, "DANH MỤC HÌNH ẢNH")
    figures_list = [
        ("Hình 1.1", "Sơ đồ chuyển đổi từ Mô hình Phân tích sang Mô hình Thiết kế trong kỹ nghệ phần mềm"),
        ("Hình 1.2", "Mô hình kiến trúc 4+1 Views của Kruchten"),
        ("Hình 1.3", "Cấu trúc cơ bản của mẫu kiến trúc Model - View - Controller (MVC)"),
        ("Hình 2.1", "Sơ đồ Use Case tổng quan của Hệ thống Quản lý và Bán Phụ tùng Xe máy (MotoParts)"),
        ("Hình 2.2", "Sơ đồ quy trình nghiệp vụ Đặt mua phụ tùng và Thanh toán trực tuyến"),
        ("Hình 2.3", "Sơ đồ quy trình nghiệp vụ Quản lý danh mục và Nhập xuất tồn kho phụ tùng xe máy"),
        ("Hình 2.4", "Sơ đồ quy trình Tư vấn kỹ thuật phụ tùng (LiveChat) và Xác thực Email OTP"),
        ("Hình 3.1", "Sơ đồ kiến trúc tổng thể 3 tầng (Layered Architecture) của hệ thống MotoParts"),
        ("Hình 3.2", "Sơ đồ phân rã thành phần hệ thống (Component Diagram)"),
        ("Hình 4.1", "Sơ đồ quan hệ thực thể (ERD) hoàn chỉnh của cơ sở dữ liệu hệ thống phụ tùng"),
        ("Hình 4.2", "Sơ đồ lớp chi tiết (Class Diagram) thể hiện mối quan hệ giữa Controllers, Models và Services"),
        ("Hình 4.3", "Sơ đồ trạng thái vòng đời Đơn hàng phụ tùng xe máy (Order State Diagram)"),
        ("Hình 4.4", "Sơ đồ trạng thái Giao dịch thanh toán MoMo"),
        ("Hình 4.5", "Sơ đồ trạng thái Tài khoản người dùng và xác thực OTP"),
        ("Hình 5.1", "Sơ đồ cấu trúc điều hướng (Sitemap & Navigation Flow) của hệ thống phụ tùng"),
        ("Hình 5.2", "Sơ đồ tuần tự: Luồng tra cứu, chọn mua phụ tùng và thanh toán MoMo"),
        ("Hình 5.3", "Sơ đồ tuần tự: Luồng tính cước phí và tích hợp tạo đơn vận chuyển Giao Hàng Nhanh (GHN)"),
        ("Hình 5.4", "Sơ đồ tuần tự: Luồng tư vấn kỹ thuật phụ tùng trực tuyến realtime giữa Khách và Kỹ thuật viên")
    ]
    add_table_styled(doc, ["Ký hiệu hình", "Tên hình ảnh và nội dung mô tả"], figures_list, [Inches(1.5), Inches(4.7)])
    doc.add_page_break()

def add_introduction(doc):
    add_heading_1(doc, "MỞ ĐẦU")
    
    add_heading_2(doc, "1. Lý do chọn đề tài")
    add_p(doc,
        "Tại Việt Nam, xe máy từ lâu đã trở thành phương tiện giao thông chủ lực, huyết mạch phục vụ việc đi lại, sinh hoạt "
        "và mưu sinh của đại đa số người dân, với tổng số lượng phương tiện lưu hành ước tính vượt trên 72 triệu xe. "
        "Đi đôi với mật độ sử dụng khổng lồ đó là nhu cầu bảo dưỡng, sửa chữa và thay thế định kỳ các loại phụ tùng, linh kiện xe máy "
        "(như bộ má phanh, lọc gió, bugi, nhông sên dĩa, hệ thống chiếu sáng đèn xi-nhan, gương chiếu hậu, dầu nhớt và dây curoa) "
        "luôn duy trì ở mức cực kỳ cao và tăng trưởng đều đặn qua các năm."
    )
    add_p(doc,
        "Tuy nhiên, phương thức kinh doanh phụ tùng truyền thống tại các cửa hàng, tiệm sửa xe vật lý hiện đang bộc lộ "
        "rất nhiều hạn chế lớn gây bức xúc cho người tiêu dùng và gây khó khăn cho công tác quản trị của chủ cửa hàng:\n"
        "1. Vấn nạn phụ tùng giả, linh kiện nhái kém chất lượng tràn lan: Người tiêu dùng thiếu công cụ tra cứu nguồn gốc xuất xứ, "
        "không nắm được mã phụ tùng chuẩn của nhà sản xuất (Mã phụ tùng OEM của Honda, Yamaha, SYM, Suzuki), dẫn đến việc dễ bị ép giá "
        "hoặc lắp đặt phụ tùng không đồng bộ, gây nguy hiểm trực tiếp đến an toàn tính mạng khi tham gia giao thông.\n"
        "2. Khó khăn trong việc tra cứu độ tương thích của phụ tùng: Khách hàng khi cần tự bảo dưỡng tại nhà rất khó biết bộ má phanh hay lọc gió nào "
        "thực sự lắp vừa cho dòng xe cụ thể của mình (ví dụ: má phanh Air Blade 125, lọc gió Wave RSX hay bugi NGK chân dài CPR6EA-9).\n"
        "3. Phương thức thanh toán và giao nhận manh mún: Đa phần các tiệm sửa xe truyền thống vẫn thu tiền mặt thủ công, chưa kết nối "
        "với các đơn vị giao hàng toàn quốc và các cổng thanh toán điện tử thông minh.\n"
        "4. Quản lý kho thủ công: Chủ cửa hàng phụ tùng gặp muôn vàn khó khăn khi theo dõi số lượng tồn kho của hàng nghìn mã phụ tùng nhỏ lẻ, "
        "dễ thất thoát hàng hóa, sai lệch số liệu doanh thu và không nắm bắt được xu hướng tiêu thụ của thị trường."
    )
    add_p(doc,
        "Dưới góc độ kỹ thuật công nghệ thông tin, việc xây dựng một hệ thống website thương mại điện tử chuyên ngành phụ tùng xe máy "
        "đòi hỏi sự kết hợp phức tạp giữa quản trị kho hàng (Inventory Management), xử lý giỏ hàng và đặt hàng đa sản phẩm, "
        "tích hợp API tính cước vận chuyển khoảng cách địa lý theo thời gian thực (qua đối tác Giao Hàng Nhanh - GHN), "
        "tích hợp Cổng thanh toán tài chính điện tử bảo mật (Ví MoMo) với chuẩn mã hóa chữ ký HMAC-SHA256, kênh chat tư vấn kỹ thuật trực tuyến (LiveChat), "
        "và hệ thống bảo mật tài khoản người dùng bằng mã OTP gửi qua Email. "
        "Để hệ thống đạt độ tin cậy cao, khả năng mở rộng linh hoạt và bảo trì dễ dàng, việc nghiên cứu chuyên sâu về Kiến trúc phần mềm "
        "(Software Architecture), tuân thủ 5 nguyên lý thiết kế hướng đối tượng SOLID và áp dụng các mẫu thiết kế GoF chuẩn mực "
        "là yếu tố then chốt quyết định sự thành công của đề tài. "
        "Chính vì những lý do thực tiễn và khoa học cấp thiết đó, nhóm nghiên cứu đã lựa chọn đề tài: "
        "'Nghiên cứu kiến trúc, nguyên lý thiết kế và xây dựng Hệ thống Thương mại Điện tử và Quản lý Cửa hàng Phụ tùng Xe máy Trực tuyến (MotoParts)' "
        "làm nội dung cho bài tập lớn của học phần Kiến trúc và Thiết kế Phần mềm."
    )
    
    add_heading_2(doc, "2. Mục tiêu nghiên cứu")
    add_p(doc, "Đề tài được xây dựng nhằm giải quyết toàn diện các mục tiêu cụ thể sau:")
    add_bullet(doc, "Hệ thống hóa toàn diện cơ sở lý luận về kỹ nghệ thiết kế phần mềm, vai trò của kiến trúc phần mềm trong chu trình phát triển (SDLC), mô hình 4+1 Views, các thuộc tính chất lượng phần mềm (FURPS+), các nguyên lý thiết kế hướng đối tượng SOLID và các mẫu thiết kế GoF.", "Về mặt lý thuyết: ");
    add_bullet(doc, "Khảo sát và phân tích toàn diện yêu cầu bài toán kinh doanh phụ tùng xe máy; xác định danh mục 15 yêu cầu chức năng (FR01 - FR15), xây dựng sơ đồ Use Case phân quyền và đặc tả chi tiết 4 Use Case cốt lõi; mô hình hóa 3 quy trình nghiệp vụ trọng tâm (Đặt hàng & Thanh toán MoMo, Quản lý kho phụ tùng, Tư vấn kỹ thuật LiveChat và Xác thực OTP).", "Về mặt phân tích yêu cầu: ");
    add_bullet(doc, "Thiết kế kiến trúc hệ thống 3 tầng (Layered Architecture) kết hợp mô hình MVC trên nền tảng Laravel Framework; phân rã hệ thống thành 8 thành phần chuyên biệt (Auth, PartsCatalog, Cart, Order, MomoPayment, GHNShipping, LiveChat, InventoryReport); thực hiện đánh giá kiến trúc theo phương pháp chuẩn quốc tế ATAM.", "Về mặt thiết kế kiến trúc: ");
    add_bullet(doc, "Thiết kế mô hình cơ sở dữ liệu quan hệ chuẩn hóa 3NF gồm 7 bảng thực thể nòng cốt; xây dựng từ điển dữ liệu chi tiết; thiết kế cấu trúc lớp hướng đối tượng (Class Diagram) và mô hình hóa động qua 3 sơ đồ máy trạng thái phản ánh chu kỳ của Đơn hàng phụ tùng, Giao dịch MoMo và Tài khoản khách hàng.", "Về mặt thiết kế dữ liệu & lớp: ");
    add_bullet(doc, "Thiết kế giao diện người dùng theo 10 nguyên lý Usability Heuristics của Jakob Nielsen; chuẩn hóa giao diện tích hợp API ngoại vi (Cổng MoMo, Giao Hàng Nhanh GHN, Máy chủ SMTP); xây dựng 3 sơ đồ tuần tự (Sequence Diagrams) và hiện thực hóa kiểm thử tự động hệ thống đảm bảo vận hành ổn định, chính xác.", "Về mặt thiết kế giao diện & thành phần: ");
    
    add_heading_2(doc, "3. Đối tượng và phạm vi nghiên cứu")
    add_p(doc, "Đối tượng nghiên cứu của đề tài bao gồm:", bold=True)
    add_bullet(doc, "Các nguyên lý, lý thuyết, tiêu chuẩn và phương pháp luận về kiến trúc phần mềm, quy trình chuyển dịch từ mô hình phân tích sang mô hình thiết kế, các mẫu kiến trúc (Architectural Patterns) và mẫu thiết kế (Design Patterns).");
    add_bullet(doc, "Nghiệp vụ thương mại điện tử chuyên ngành linh kiện, phụ tùng xe máy: phân loại danh mục phụ tùng, mã OEM, quản lý tồn kho, giỏ hàng, cước vận chuyển bưu chính và thanh toán ví điện tử.");
    add_bullet(doc, "Framework kiến trúc ứng dụng web Laravel (PHP 8.2+), hệ quản trị cơ sở dữ liệu quan hệ MySQL 8.0, chuẩn bảo mật HMAC-SHA256, API RESTful của MoMo và GHN.");
    
    add_p(doc, "Phạm vi nghiên cứu và ứng dụng của đề tài:", bold=True)
    add_bullet(doc, "Về mặt không gian nghiệp vụ: Hệ thống phục vụ trọn vẹn chu trình mua sắm phụ tùng của khách hàng (tra cứu mã phụ tùng, lọc theo danh mục, thêm giỏ hàng, tính phí ship GHN, thanh toán qua Ví MoMo hoặc COD) và chu trình quản trị của chủ cửa hàng (quản lý sản phẩm, cập nhật tồn kho, duyệt đơn hàng, theo dõi trạng thái giao vận, chat tư vấn kỹ thuật và xem báo cáo doanh thu).");
    add_bullet(doc, "Về mặt kỹ thuật: Xây dựng hoàn chỉnh ứng dụng web theo phong cách MVC phân tầng trên Laravel, xác thực phân quyền qua Middleware (Guest, Customer, Admin), xác thực email hai bước qua mã OTP 6 số, bảo vệ chống giả mạo dữ liệu.");
    add_bullet(doc, "Giới hạn đề tài: Không nghiên cứu các máy móc cơ khí đo kiểm thông số động cơ vật lý tại xưởng sửa chữa hay hệ thống kế toán doanh nghiệp đa chi nhánh ERP quy mô tập đoàn.");

    add_heading_2(doc, "4. Phương pháp thực hiện")
    add_p(doc, "Đề tài áp dụng phối hợp các phương pháp nghiên cứu khoa học và kỹ nghệ phần mềm chuẩn mực sau:")
    add_bullet(doc, "Thu thập và phân tích tài liệu chuẩn mực về kỹ nghệ phần mềm của Roger S. Pressman, Ian Sommerville, Len Bass, Paul Clements, Rick Kazman và Robert C. Martin.", "Phương pháp nghiên cứu lý thuyết: ");
    add_bullet(doc, "Sử dụng ngôn ngữ mô hình hóa thống nhất UML (Use Case, Activity, Component, Class, Sequence, State Machine Diagrams) để trực quan hóa kiến trúc hệ thống.", "Phương pháp mô hình hóa hướng đối tượng: ");
    add_bullet(doc, "Áp dụng phương pháp phân tích đánh đổi kiến trúc (ATAM) của Viện SEI để kiểm chứng các kịch bản chất lượng khắt khe về Hiệu năng, Độ tin cậy và Bảo mật.", "Phương pháp đánh giá kiến trúc: ");
    add_bullet(doc, "Cài đặt ứng dụng thực tế trên Laravel/MySQL, viết bộ kịch bản kiểm thử tự động (Feature Testing) với PHPUnit để kiểm tra độ tin cậy của mã nguồn.", "Phương pháp thực nghiệm & kiểm thử: ");

    add_heading_2(doc, "5. Cấu trúc báo cáo")
    add_p(doc, "Báo cáo bài tập lớn được kết cấu thành 5 chương chính cùng phần mở đầu, kết luận và tài liệu tham khảo theo đúng đề cương chuẩn mực:")
    add_bullet(doc, "Trang bìa chính, Trang bìa phụ, Lời cảm ơn, Mục lục, Danh mục chữ viết tắt, Danh mục bảng biểu, Danh mục hình ảnh.", "Phần đầu: ");
    add_bullet(doc, "Trình bày lý do chọn đề tài, mục tiêu, đối tượng, phạm vi, phương pháp và bố cục báo cáo.", "Phần Mở đầu: ");
    add_bullet(doc, "Trình bày toàn diện cơ sở lý thuyết về thiết kế phần mềm trong SDLC, chuyển đổi sang mô hình thiết kế, kiến trúc phần mềm, thuộc tính chất lượng, yêu cầu có ý nghĩa kiến trúc (ASRs), 4 nhóm nguyên lý thiết kế (SOLID), thể loại & phong cách kiến trúc, mẫu thiết kế GoF.", "Chương 1 - Tổng quan về Kiến trúc và Thiết kế phần mềm: ");
    add_bullet(doc, "Tổng quan bài toán cửa hàng phụ tùng xe máy, danh mục 15 yêu cầu chức năng, sơ đồ Use Case, đặc tả 4 Use Case trọng tâm, mô hình hóa 3 quy trình nghiệp vụ nòng cốt, yêu cầu phi chức năng (FURPS+) và các ràng buộc.", "Chương 2 - Phân tích yêu cầu và Thiết kế mô hình: ");
    add_bullet(doc, "Tiêu chí lựa chọn, ma trận quyết định chọn kiến trúc phân tầng MVC trên Laravel, kiến trúc tổng thể 3 tầng, phân rã 8 thành phần, sơ đồ Component Diagram và đánh giá kiến trúc theo phương pháp ATAM.", "Chương 3 - Thiết kế kiến trúc: ");
    add_bullet(doc, "Thiết kế CSDL quan hệ 3NF gồm 7 bảng thực thể (users, categories, products, orders, order_items, payment_transactions, messages), từ điển dữ liệu, sơ đồ ERD, sơ đồ lớp (Class Diagram), và 3 sơ đồ máy trạng thái cho Đơn hàng, Giao dịch MoMo, Tài khoản người dùng.", "Chương 4 - Thiết kế dữ liệu và lớp: ");
    add_bullet(doc, "Thiết kế giao diện theo 10 nguyên lý Nielsen Heuristics, sơ đồ Sitemap điều hướng, đặc tả các màn hình chính; thiết kế tích hợp API bên ngoài (MoMo, GHN, SMTP); thiết kế thành phần, interface dịch vụ, 3 sơ đồ tuần tự (Sequence Diagrams) và đánh giá chất lượng thành phần.", "Chương 5 - Thiết kế giao diện và thành phần: ");
    add_bullet(doc, "Tổng kết các kết quả đạt được, chỉ ra các hạn chế và đề xuất hướng mở rộng nâng cấp hệ thống trong tương lai.", "Phần Kết luận và Hướng phát triển: ");
    add_bullet(doc, "Danh mục các tài liệu tham khảo khoa học uy tín trích dẫn theo chuẩn mực học thuật quốc tế.", "Tài liệu tham khảo: ");
    
    doc.add_page_break()
