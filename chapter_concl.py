import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_styles import (
    add_p, add_heading_1, add_heading_2, add_heading_3, add_bullet, 
    COLOR_PRIMARY, COLOR_SECONDARY, COLOR_TEXT
)

def build_conclusion_and_references(doc):
    add_heading_1(doc, "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    
    add_heading_2(doc, "1. Kết luận")
    add_p(doc,
        "Đề tài bài tập lớn môn học Kiến trúc và Thiết kế Phần mềm với tên đề tài: 'Nghiên cứu kiến trúc, nguyên lý thiết kế "
        "và xây dựng Hệ thống Thương mại Điện tử và Quản lý Cửa hàng Phụ tùng Xe máy Trực tuyến (MotoParts)' "
        "đã được nhóm nghiên cứu triển khai một cách nghiêm túc, khoa học, bám sát các phương pháp luận kỹ nghệ phần mềm hiện đại "
        "và đạt được những kết quả đáng ghi nhận."
    )
    
    add_heading_3(doc, "1.1. Kết quả đạt được")
    add_bullet(doc, "Đã hệ thống hóa toàn diện cơ sở lý thuyết về thiết kế phần mềm trong chu trình phát triển (SDLC), vai trò của kiến trúc phần mềm, mô hình kiến trúc 4+1 Views, các thuộc tính chất lượng (FURPS+), nguyên lý phân rã mô đun hóa (Cohesion & Coupling), bộ 5 nguyên lý hướng đối tượng SOLID và 23 mẫu thiết kế GoF kinh điển.", "Về mặt học thuật và lý thuyết: ");
    add_bullet(doc, "Đã khảo sát sâu sắc thực trạng bài toán kinh doanh phụ tùng xe máy tại Việt Nam; xác định tường minh 15 yêu cầu chức năng, sơ đồ Use Case phân quyền rành mạch (Guest, Customer, Admin); đặc tả chi tiết 4 Use Case cốt lõi và mô hình hóa trực quan 3 quy trình nghiệp vụ then chốt (Đặt hàng & MoMo, Quản lý kho phụ tùng, Tư vấn kỹ thuật LiveChat & Xác thực OTP).", "Về mặt phân tích yêu cầu: ");
    add_bullet(doc, "Đã thiết lập ma trận quyết định khoa học để lựa chọn phong cách kiến trúc Phân tầng kết hợp mô hình MVC (Modular Layered Monolith) trên Laravel Framework; phân rã thành công hệ thống thành 8 thành phần chuyên biệt; đánh giá kiến trúc theo phương pháp chuẩn quốc tế ATAM qua các kịch bản chất lượng khắt khe.", "Về mặt thiết kế kiến trúc: ");
    add_bullet(doc, "Đã thiết kế cơ sở dữ liệu quan hệ chuẩn hóa 3NF gồm 7 bảng thực thể nòng cốt (`users`, `categories`, `products`, `orders`, `order_items`, `payment_transactions`, `messages`), đảm bảo tính toàn vẹn tham chiếu ACID; thiết kế cấu trúc lớp tĩnh (Class Diagram) và mô hình hóa hành vi động qua 3 sơ đồ máy trạng thái (Statechart Diagrams).", "Về mặt thiết kế dữ liệu và lớp: ");
    add_bullet(doc, "Đã thiết kế giao diện người dùng đạt chuẩn 10 nguyên lý Usability Heuristics của Jakob Nielsen; chuẩn hóa giao diện tích hợp API với bên ngoài (Cổng MoMo HMAC-SHA256, Giao Hàng Nhanh GHN, SMTP); xây dựng 3 sơ đồ tuần tự (Sequence Diagrams) và hiện thực hóa kiểm thử tự động xác nhận hệ thống vận hành hoàn hảo, không có lỗi.", "Về mặt thiết kế giao diện và thành phần: ");

    add_heading_3(doc, "1.2. Hạn chế")
    add_p(doc, "Bên cạnh các kết quả xuất sắc đã đạt được, hệ thống vẫn còn tồn tại một số điểm hạn chế do giới hạn về mặt thời gian và nguồn lực:")
    add_bullet(doc, "Hệ thống hiện tại chưa tích hợp tính năng quét mã vạch (Barcode / QR Code) tại quầy để hỗ trợ nhân viên kho nhập và xuất phụ tùng vật lý nhanh chóng bằng máy quét chuyên dụng.");
    add_bullet(doc, "Cơ chế trao đổi tin nhắn LiveChat hiện sử dụng cơ chế Polling / SSE thông qua AJAX định kỳ, chưa nâng cấp lên giao thức WebSocket hai chiều thực sự (như Laravel Echo kết hợp Pusher/Soketi) để tối ưu hóa triệt để tài nguyên máy chủ khi có lượng truy cập lớn.");
    add_bullet(doc, "Chưa tích hợp công cụ AI nhận diện hình ảnh phụ tùng hoặc gợi ý phụ tùng tự động dựa trên ảnh chụp thực tế linh kiện bị hỏng của khách hàng.");

    add_heading_2(doc, "2. Hướng phát triển")
    add_p(doc, "Trong thời gian tới, nhóm nghiên cứu đề xuất các định hướng nâng cấp và mở rộng hệ thống như sau:")
    add_bullet(doc, "Xây dựng tính năng 'Tra cứu phụ tùng thông minh theo dòng xe': Khách hàng chỉ cần chọn Hãng xe (Honda, Yamaha, Suzuki) -> Dòng xe (Air Blade, Wave, Exciter) -> Năm sản xuất (2018, 2020, 2024), hệ thống sẽ tự động lọc ra toàn bộ phụ tùng tương thích 100%.");
    add_bullet(doc, "Nâng cấp giao tiếp LiveChat và Thông báo trạng thái vận chuyển GHN theo thời gian thực lên nền tảng WebSocket (Laravel Reverb) nhằm giảm độ trễ trao đổi kỹ thuật xuống mức mili-giây.");
    add_bullet(doc, "Mở rộng tích hợp thêm các cổng thanh toán tài chính phổ biến khác như VNPay, ZaloPay, ShopeePay và tính năng tạo mã VietQR động tự động điền số tiền và nội dung chuyển khoản.");
    add_bullet(doc, "Tích hợp ứng dụng quét mã vạch trên điện thoại di động giúp nhân viên kho kiểm kê số lượng tồn kho phụ tùng nhanh chóng, tự động cảnh báo khi một mã phụ tùng chạm ngưỡng tồn kho tối thiểu.");
    add_bullet(doc, "Xây dựng ứng dụng di động đa nền tảng (Mobile App trên Flutter / React Native) dành riêng cho thợ sửa xe và khách hàng thân thiết, tích hợp tính năng đặt lịch sửa xe và thay thế phụ tùng tận nhà.");

    doc.add_page_break()

    # -------------------------------------------------------------
    # TÀI LIỆU THAM KHẢO
    # -------------------------------------------------------------
    add_heading_1(doc, "TÀI LIỆU THAM KHẢO")
    
    references = [
        "[1] Roger S. Pressman, Ph.D., Bruce R. Maxim, Ph.D. (2020), Software Engineering: A Practitioner's Approach, 9th Edition, McGraw-Hill Education.",
        "[2] Ian Sommerville (2016), Software Engineering, 10th Edition, Pearson Education.",
        "[3] Len Bass, Paul Clements, Rick Kazman (2021), Software Architecture in Practice, 4th Edition, Addison-Wesley Professional (SEI Series in Software Engineering).",
        "[4] Robert C. Martin (Uncle Bob) (2018), Clean Architecture: A Craftsman's Guide to Software Structure and Design, Prentice Hall.",
        "[5] Robert C. Martin (2002), Agile Software Development, Principles, Patterns, and Practices, Pearson.",
        "[6] Erich Gamma, Richard Helm, Ralph Johnson, John Vlissides (Gang of Four - GoF) (1994), Design Patterns: Elements of Reusable Object-Oriented Software, Addison-Wesley Professional.",
        "[7] Martin Fowler (2002), Patterns of Enterprise Application Architecture, Addison-Wesley Professional.",
        "[8] Philippe Kruchten (1995), 'Architectural Blueprints — The '4+1' View Model of Software Architecture', IEEE Software, vol. 12, no. 6, pp. 42-50.",
        "[9] Jakob Nielsen (1994), Usability Engineering, Morgan Kaufmann Publishers.",
        "[10] Paul Clements, Rick Kazman, Mark Klein (2002), Evaluating Software Architectures: Methods and Case Studies, Addison-Wesley Professional.",
        "[11] Taylor Otwell (2024), Laravel Documentation (The PHP Framework for Web Artisans), https://laravel.com/docs.",
        "[12] MoMo Payment Platform (2024), MoMo Payment API Developer Documentation & HMAC-SHA256 Guidelines, https://developers.momo.vn.",
        "[13] Giao Hàng Nhanh (GHN) (2024), GHN Express API Open Documentation, https://api.ghn.vn."
    ]
    
    for ref in references:
        add_p(doc, ref, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
